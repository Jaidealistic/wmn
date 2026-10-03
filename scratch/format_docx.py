import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import sys

def format_document(input_path, output_path):
    try:
        doc = docx.Document(input_path)
    except Exception as e:
        print(f"Error opening {input_path}: {e}")
        return

    # Figure titles in order
    captions = [
        "Figure 1: Design 1 - Normal Base Topology (Node 1 to Node 15)",
        "Figure 2: Design 2 - Link Failure (Distance gap between Node 7 and 8)",
        "Figure 3: Design 3 - Node 8 Failure (Powered off/Dead)",
        "Figure 4: Design 4 - Both Link and Node Failures"
    ]
    image_count = 0

    # We will iterate through paragraphs to find the placeholder and images
    paragraphs = doc.paragraphs
    for i, p in enumerate(paragraphs):
        # Remove placeholder text
        if "IMPORTANT: I CANNOT AUTOMATICALLY PASTE" in p.text:
            p.text = "" # Clear it out
        
        # Check if the paragraph contains an image
        has_image = False
        for run in p.runs:
            if '<w:drawing' in run._element.xml:
                has_image = True
                break
        
        if has_image:
            # Center the image
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Insert a caption after this image
            if image_count < len(captions):
                # We insert before the next paragraph, which effectively puts it after the current one
                if i + 1 < len(paragraphs):
                    new_p = paragraphs[i+1].insert_paragraph_before(captions[image_count])
                else:
                    new_p = doc.add_paragraph(captions[image_count])
                
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                new_r = new_p.runs[0]
                new_r.font.name = 'Arial'
                new_r.font.size = Pt(10)
                new_r.font.italic = True
                
                # Add some spacing after the caption
                if i + 1 < len(paragraphs):
                    paragraphs[i+1].insert_paragraph_before("")

                image_count += 1

    try:
        doc.save(output_path)
        print(f"Success! Formatted document saved to {output_path}")
        print(f"Found and captioned {image_count} images.")
    except Exception as e:
        print(f"Error saving document: {e}")

if __name__ == "__main__":
    format_document(r"c:\Users\jaisu\Projects\wmn\cs2.docx", r"c:\Users\jaisu\Projects\wmn\CaseStudy2_Submission_Ready.docx")
